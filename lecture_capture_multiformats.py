# =============================================================================
#  LECTURE CAPTURE
#  Take screenshots of a lecture with a keyboard shortcut and turn them into
#  an ordered PowerPoint, PDF or Word file (or a ZIP of the images).
# =============================================================================
#
#  HOW TO INSTALL AND RUN  (Windows 10 / 11)
#  -----------------------------------------
#  STEP 1  Install Python (free) from  https://www.python.org/downloads/
#          IMPORTANT: on the first installer screen, tick the box
#          "Add python.exe to PATH" before clicking Install.
#
#  STEP 2  Save this file (lecture_capture.py) in any folder, for example
#          a new folder on your Desktop.
#
#  STEP 3  Open that folder, click on the address bar at the top, type  cmd
#          and press Enter. A black Command Prompt window opens.
#
#  STEP 4  Install the helper libraries (only needed ONE time).
#          Copy this line, paste it into the black window and press Enter:
#
#              python -m pip install customtkinter mss keyboard python-pptx python-docx pillow
#
#  STEP 5  Start the app:
#
#              python lecture_capture.py
#
#  If Windows says "python is not recognized", type  py  instead of  python
#  in steps 4 and 5.
#
#  HOW TO USE IT
#  -------------
#   1. Type a lecture name, pick your monitor, press "Start Lecture".
#   2. Watch your lecture. Press the capture key (default F8) whenever you
#      want a screenshot. Pressed twice by mistake? Press F9 to delete it.
#   3. Tick the export formats you want (PowerPoint, PDF, Word, or a ZIP of
#      the images). Press "Stop & Export", choose a name and folder, done.
#      Every ticked format is saved with that name, screenshots in order.
#      The temporary screenshots are deleted automatically after saving.
#
#  OPTIONAL: MAKE A DOUBLE-CLICKABLE .EXE (so you don't need the black window)
#          python -m pip install pyinstaller
#          pyinstaller --onefile --noconsole lecture_capture.py
#          The .exe appears in a new folder called "dist".
#
#  PROBLEMS?
#   - Shortcut keys do nothing: close the app, open Command Prompt by
#     right-clicking it and choosing "Run as administrator", then start again.
#   - Screenshots are black: turn off "hardware acceleration" in your browser
#     or video player, then try again.
#   - Shortcut clashes with another program (like Ctrl+Shift+S in Firefox):
#     click the key button in the app and press a different combination.
#   - This app works on Windows only.
# =============================================================================


# ---------------------------------------------------------------------------
# 1. IMPORTS - tools (libraries) our program uses
# ---------------------------------------------------------------------------
import json        # save/load our settings in a small text file
import shutil      # delete whole folders
import subprocess  # open Windows Explorer at the end
import threading   # wait for a key press without freezing the window
import time        # small pauses
import zipfile     # make .zip files
import tkinter as tk                          # basic window tools (used for the region picker)
from datetime import datetime                 # current date and time
from pathlib import Path                      # easy file and folder paths
from tkinter import filedialog, messagebox    # "Save as" window and pop-up questions

import customtkinter as ctk   # modern looking buttons, cards and switches
import keyboard               # listens for shortcut keys, even in other apps
import mss                    # takes the screenshots
from PIL import Image, ImageChops, ImageStat   # edit/compare pictures
from pptx import Presentation                  # builds the PowerPoint file
from pptx.util import Inches

# Word support is optional: if python-docx is not installed the app still
# starts, and only the "Word" option shows a reminder.
try:
    from docx import Document                      # builds Word files
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches as DocInches
except ImportError:
    Document = None

try:
    import winsound           # makes the "beep" sound (Windows only)
except ImportError:
    winsound = None

# Newer versions of "mss" call the screenshot tool MSS, older ones call it mss.
# This line picks whichever one exists so the app works with both.
MSS = getattr(mss, "MSS", None) or mss.mss


# ---------------------------------------------------------------------------
# 2. SETTINGS AND COLOURS - change these to customise the app
# ---------------------------------------------------------------------------
# Where temporary files and settings live: Documents/LectureCaptures
APP_FOLDER = Path.home() / "Documents" / "LectureCaptures"
TEMP_FOLDER = APP_FOLDER / "_in_progress"      # screenshots wait here until export
SETTINGS_FILE = APP_FOLDER / "settings.json"

# Values used the first time you run the app (they are remembered afterwards)
DEFAULT_SETTINGS = {
    "capture_key": "f8",                 # shortcut to take a screenshot
    "undo_key": "f9",                    # shortcut to delete the last screenshot
    "monitor": 1,                        # which monitor to capture
    "note_slides": False,                # add a blank slide after each screenshot?
    "fmt_pptx": True,                    # export as PowerPoint?
    "fmt_pdf": False,                    # export as PDF?
    "fmt_docx": False,                   # export as Word document?
    "fmt_zip": False,                    # export a ZIP of the images?
    "skip_duplicates": False,            # skip a screenshot that looks the same as the last?
    "beep": True,                        # play a beep after each screenshot?
    "last_folder": str(Path.home() / "Documents"),   # last place you saved to
}

# PowerPoint slide size in inches (widescreen 16:9)
SLIDE_WIDTH, SLIDE_HEIGHT = 13.333, 7.5

# Export formats: short name -> (label shown in the app, file ending)
FORMATS = {
    "pptx": ("PowerPoint", ".pptx"),
    "pdf": ("PDF", ".pdf"),
    "docx": ("Word", ".docx"),
    "zip": ("Images (ZIP)", ".zip"),
}

# Colours (hex codes). Change them to give the app a different look!
BG = "#0d0f1f"        # window background (deep navy)
CARD = "#171a33"      # card background
FIELD = "#232850"     # text box / key button background
VIOLET = "#7c5cff"
CYAN = "#22d3ee"
PINK = "#ff4d8d"
GREEN = "#22c55e"
ORANGE = "#f59e0b"
TEXT = "#e8eaff"      # main text
MUTED = "#8b90b8"     # soft grey-blue text


# ---------------------------------------------------------------------------
# 3. HELPER FUNCTIONS - small jobs that do not need the window
# ---------------------------------------------------------------------------
def load_settings():
    """Read settings.json. If it does not exist yet, use the defaults."""
    settings = dict(DEFAULT_SETTINGS)
    try:
        saved = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        settings.update(saved)
    except Exception:
        pass   # no file yet (first run) - that is fine
    return settings


def save_settings(settings):
    """Write the settings to settings.json so the app remembers them."""
    try:
        APP_FOLDER.mkdir(parents=True, exist_ok=True)
        SETTINGS_FILE.write_text(json.dumps(settings, indent=2), encoding="utf-8")
    except Exception:
        pass


def clean_file_name(text):
    """Remove characters Windows does not allow in file names."""
    return "".join(c for c in text if c.isalnum() or c in " -_").strip()


def images_look_same(image_a, image_b, threshold=0.5):
    """Return True if two screenshots are almost identical.
    We shrink both to tiny grey pictures and measure the average difference."""
    small_a = image_a.convert("L").resize((96, 54))
    small_b = image_b.convert("L").resize((96, 54))
    difference = ImageChops.difference(small_a, small_b)
    return ImageStat.Stat(difference).mean[0] < threshold


def build_pptx(image_paths, output_path, add_note_slides=False):
    """Make a PowerPoint: one screenshot per slide, in the order of the list."""
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH)
    presentation.slide_height = Inches(SLIDE_HEIGHT)
    blank_layout = presentation.slide_layouts[6]    # layout number 6 = empty slide

    for path in image_paths:
        slide = presentation.slides.add_slide(blank_layout)

        # Find the picture size so we can fit it on the slide without stretching
        with Image.open(path) as picture:
            width, height = picture.size
        scale = min(SLIDE_WIDTH / width, SLIDE_HEIGHT / height)
        new_width, new_height = width * scale, height * scale

        # Put the picture in the middle of the slide
        left = (SLIDE_WIDTH - new_width) / 2
        top = (SLIDE_HEIGHT - new_height) / 2
        slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                                 Inches(new_width), Inches(new_height))

        if add_note_slides:
            presentation.slides.add_slide(blank_layout)   # empty slide for your notes

    presentation.save(str(output_path))


def build_pdf(image_paths, output_path, add_note_pages=False):
    """Make a PDF with one screenshot per page (optionally a blank page after each)."""
    pages = []
    for path in image_paths:
        picture = Image.open(path).convert("RGB")
        pages.append(picture)
        if add_note_pages:
            pages.append(Image.new("RGB", picture.size, "white"))   # empty page for notes
    pages[0].save(str(output_path), save_all=True, append_images=pages[1:])


def build_docx(image_paths, output_path, add_note_pages=False):
    """Make a Word document: one screenshot per landscape page, in order."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = DocInches(11), DocInches(8.5)
    section.left_margin = section.right_margin = DocInches(0.5)
    section.top_margin = section.bottom_margin = DocInches(0.5)

    max_width, max_height = 10.0, 6.9      # biggest picture that fits on the page
    for number, path in enumerate(image_paths):
        with Image.open(path) as picture:
            width, height = picture.size
        scale = min(max_width / width, max_height / height)
        document.add_picture(str(path), width=DocInches(width * scale),
                             height=DocInches(height * scale))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        if add_note_pages:                 # a page headed "Notes" for your own writing
            document.add_page_break()
            document.add_paragraph("Notes")
        if number < len(image_paths) - 1:
            document.add_page_break()
    document.save(str(output_path))


def build_zip(image_paths, output_path):
    """Put all the screenshots (001.png, 002.png, ...) into one .zip file."""
    with zipfile.ZipFile(str(output_path), "w", zipfile.ZIP_STORED) as archive:
        for path in image_paths:
            archive.write(str(path), arcname=Path(path).name)


# ---------------------------------------------------------------------------
# 4. THE APP WINDOW
#    "self" below simply means "this app window". Things stored on self
#    (like self.screenshots) can be used by every function in the class.
# ---------------------------------------------------------------------------
class App(ctk.CTk):

    # ---- 4a. Setting everything up when the app starts ---------------------
    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("dark")
        self.title("Lecture Capture")
        self.geometry("800x680")
        self.resizable(False, False)
        self.attributes("-topmost", True)       # keep the window above other windows
        self.configure(fg_color=BG)

        # Things the app needs to remember while running
        self.settings = load_settings()
        self.capture_key = self.settings["capture_key"]
        self.undo_key = self.settings["undo_key"]
        self.lecture_running = False       # True between Start and Stop
        self.listening_for_key = False     # True while you are choosing a new shortcut
        self.screenshots = []              # list of saved picture files, in order
        self.session_folder = None         # temporary folder for this lecture
        self.region = None                 # None = whole monitor, otherwise a selected box
        self.last_image = None             # previous screenshot (for duplicate check)
        self.hotkey_handles = []           # shortcuts currently registered
        self.blink_on = False              # used for the blinking RECORDING light

        self.build_screen()
        self.blink()                       # start the blinking light loop
        self.protocol("WM_DELETE_WINDOW", self.on_close)   # run on_close when X is clicked

    # ---- 4b. Building the screen (all the visible parts) -------------------
    def build_screen(self):
        # Header at the top
        header = ctk.CTkFrame(self, fg_color="transparent")
        header.pack(fill="x", padx=22, pady=(16, 4))
        ctk.CTkLabel(header, text="Lecture Capture", text_color=TEXT,
                     font=ctk.CTkFont("Segoe UI", 26, "bold")).pack(anchor="w")
        ctk.CTkLabel(header, text="Screenshot your lecture  ->  get ordered notes",
                     text_color=MUTED, font=ctk.CTkFont("Segoe UI", 13)).pack(anchor="w")

        # Two columns under the header: left and right
        columns = ctk.CTkFrame(self, fg_color="transparent")
        columns.pack(fill="both", expand=True, padx=12, pady=4)
        columns.grid_columnconfigure((0, 1), weight=1, uniform="cols")
        left = ctk.CTkFrame(columns, fg_color="transparent")
        right = ctk.CTkFrame(columns, fg_color="transparent")
        left.grid(row=0, column=0, sticky="nsew")
        right.grid(row=0, column=1, sticky="nsew")

        self.build_lecture_card(left)
        self.build_shortcut_card(left)
        self.build_formats_card(left)
        self.build_options_card(right)
        self.build_counter_card(right)
        self.build_buttons(right)

    def make_card(self, parent, title, title_color):
        """Create a rounded card with a coloured title and return it."""
        card = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=16)
        card.pack(fill="x", padx=8, pady=6)
        ctk.CTkLabel(card, text=title, text_color=title_color,
                     font=ctk.CTkFont("Segoe UI", 14, "bold")).pack(anchor="w", padx=16, pady=(10, 4))
        return card

    def make_button(self, parent, text, color, hover, command, height=38):
        """Create a rounded coloured button."""
        return ctk.CTkButton(parent, text=text, command=command, height=height,
                             fg_color=color, hover_color=hover, corner_radius=12,
                             text_color_disabled="#6b6f99",
                             font=ctk.CTkFont("Segoe UI", 14, "bold"))

    # -- Card 1: lecture name, monitor and capture area
    def build_lecture_card(self, parent):
        card = self.make_card(parent, "LECTURE", CYAN)

        self.name_box = ctk.CTkEntry(card, placeholder_text="Lecture name (used as the file name)",
                                     height=36, fg_color=FIELD, border_width=0, corner_radius=10)
        self.name_box.pack(fill="x", padx=16, pady=4)

        # Count monitors so the drop-down only shows the ones you have
        with MSS() as sct:
            monitor_count = len(sct.monitors) - 1     # item 0 means "all monitors together"
        self.monitor_menu = ctk.CTkOptionMenu(
            card, values=[f"Monitor {i}" for i in range(1, monitor_count + 1)],
            height=36, fg_color=FIELD, button_color=VIOLET, button_hover_color="#6747e8",
            corner_radius=10, command=self.on_monitor_changed)
        self.monitor_menu.set(f"Monitor {min(self.settings['monitor'], monitor_count)}")
        self.monitor_menu.pack(fill="x", padx=16, pady=4)

        row = ctk.CTkFrame(card, fg_color="transparent")
        row.pack(fill="x", padx=16, pady=4)
        row.grid_columnconfigure((0, 1), weight=1, uniform="r")
        self.make_button(row, "Select region", VIOLET, "#6747e8",
                         self.pick_region, 34).grid(row=0, column=0, padx=(0, 4), sticky="ew")
        self.make_button(row, "Full screen", FIELD, "#2f3566",
                         self.use_full_screen, 34).grid(row=0, column=1, padx=(4, 0), sticky="ew")

        self.area_label = ctk.CTkLabel(card, text="Capture area: full screen",
                                       text_color=MUTED, font=ctk.CTkFont("Segoe UI", 12))
        self.area_label.pack(pady=(2, 10))

    # -- Card 2: shortcut keys
    def build_shortcut_card(self, parent):
        card = self.make_card(parent, "SHORTCUT KEYS", PINK)

        grid = ctk.CTkFrame(card, fg_color="transparent")
        grid.pack(fill="x", padx=16, pady=2)
        grid.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(grid, text="Capture", text_color=TEXT).grid(row=0, column=0, sticky="w", pady=3)
        ctk.CTkLabel(grid, text="Delete last", text_color=TEXT).grid(row=1, column=0, sticky="w", pady=3)

        # These two buttons show the current shortcut. Click one, then press the
        # keys you want. "lambda" just lets us pass the word "capture"/"undo" along.
        self.capture_key_button = self.make_button(
            grid, self.capture_key, FIELD, "#2f3566", lambda: self.record_key("capture"), 34)
        self.capture_key_button.grid(row=0, column=1, padx=(12, 0), pady=3, sticky="ew")
        self.undo_key_button = self.make_button(
            grid, self.undo_key, FIELD, "#2f3566", lambda: self.record_key("undo"), 34)
        self.undo_key_button.grid(row=1, column=1, padx=(12, 0), pady=3, sticky="ew")

        self.make_button(card, "Save keys", PINK, "#e03a77",
                         self.apply_keys, 34).pack(fill="x", padx=16, pady=(6, 12))

    # -- Card 3: which file formats to export
    def build_formats_card(self, parent):
        card = self.make_card(parent, "EXPORT FORMATS", ORANGE)
        grid = ctk.CTkFrame(card, fg_color="transparent")
        grid.pack(fill="x", padx=16, pady=2)
        grid.grid_columnconfigure((0, 1), weight=1, uniform="f")

        self.format_vars = {}      # remembers which boxes are ticked
        position = 0
        for key, (label, _ending) in FORMATS.items():
            variable = ctk.BooleanVar(value=self.settings["fmt_" + key])
            self.format_vars[key] = variable
            ctk.CTkCheckBox(grid, text=label, variable=variable, text_color=TEXT,
                            fg_color=ORANGE, hover_color="#d98a06",
                            font=ctk.CTkFont("Segoe UI", 13)
                            ).grid(row=position // 2, column=position % 2, sticky="w", pady=5)
            position += 1

        ctk.CTkLabel(card, text="Tick one or more. All use the same file name.",
                     text_color=MUTED, font=ctk.CTkFont("Segoe UI", 12)
                     ).pack(anchor="w", padx=16, pady=(2, 12))

    # -- Card 4: on/off options
    def build_options_card(self, parent):
        card = self.make_card(parent, "OPTIONS", GREEN)
        self.note_var = ctk.BooleanVar(value=self.settings["note_slides"])
        self.dup_var = ctk.BooleanVar(value=self.settings["skip_duplicates"])
        self.beep_var = ctk.BooleanVar(value=self.settings["beep"])
        choices = (("Blank note page after each screenshot", self.note_var),
                   ("Skip near-identical screenshots", self.dup_var),
                   ("Beep after each screenshot", self.beep_var))
        for text, variable in choices:
            ctk.CTkSwitch(card, text=text, variable=variable, text_color=TEXT,
                          progress_color=VIOLET, font=ctk.CTkFont("Segoe UI", 13)
                          ).pack(anchor="w", padx=16, pady=4)
        ctk.CTkLabel(card, text="").pack(pady=2)   # small spacer at the bottom

    # -- Card 5: the big page counter and the blinking light
    def build_counter_card(self, parent):
        card = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=16)
        card.pack(fill="x", padx=8, pady=6)
        self.counter_label = ctk.CTkLabel(card, text="0", text_color=CYAN,
                                          font=ctk.CTkFont("Segoe UI", 52, "bold"))
        self.counter_label.pack(pady=(8, 0))
        ctk.CTkLabel(card, text="pages captured", text_color=MUTED).pack()
        self.rec_label = ctk.CTkLabel(card, text="READY", text_color=MUTED,
                                      font=ctk.CTkFont("Segoe UI", 12, "bold"))
        self.rec_label.pack(pady=(2, 10))

    # -- The four big action buttons and the status line
    def build_buttons(self, parent):
        self.start_button = self.make_button(parent, "Start Lecture", GREEN, "#16a34a",
                                             self.start_lecture, 44)
        self.start_button.pack(fill="x", padx=8, pady=(6, 3))

        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(fill="x", padx=8, pady=3)
        row.grid_columnconfigure((0, 1), weight=1, uniform="b")
        self.capture_button = self.make_button(row, f"Capture ({self.capture_key})", VIOLET,
                                               "#6747e8", self.take_screenshot, 40)
        self.capture_button.grid(row=0, column=0, padx=(0, 4), sticky="ew")
        self.undo_button = self.make_button(row, f"Delete ({self.undo_key})", ORANGE,
                                            "#d98a06", self.delete_last, 40)
        self.undo_button.grid(row=0, column=1, padx=(4, 0), sticky="ew")

        self.stop_button = self.make_button(parent, "Stop & Export", PINK, "#e03a77",
                                            self.stop_and_export, 44)
        self.stop_button.pack(fill="x", padx=8, pady=3)

        # Buttons that only make sense during a lecture start switched off
        for button in (self.capture_button, self.undo_button, self.stop_button):
            button.configure(state="disabled")

        self.status_label = ctk.CTkLabel(parent, text="Ready", text_color=MUTED,
                                         font=ctk.CTkFont("Segoe UI", 12))
        self.status_label.pack(pady=6)

    # ---- 4c. Small helpers used by many functions ---------------------------
    def set_status(self, message, color=MUTED):
        """Show a message at the bottom of the window."""
        self.status_label.configure(text=message, text_color=color)

    def update_counter(self):
        self.counter_label.configure(text=str(len(self.screenshots)))

    def selected_monitor(self):
        """Turn the text 'Monitor 2' into the number 2."""
        return int(self.monitor_menu.get().split()[-1])

    def blink(self):
        """Runs every 0.6 seconds forever: blinks 'RECORDING' during a lecture."""
        if self.lecture_running:
            self.blink_on = not self.blink_on
            self.rec_label.configure(text="RECORDING",
                                     text_color=PINK if self.blink_on else MUTED)
        else:
            self.rec_label.configure(text="READY", text_color=MUTED)
        self.after(600, self.blink)     # call this function again in 600 milliseconds

    def remember_settings(self):
        """Save the current choices so they are there next time."""
        self.settings.update(
            capture_key=self.capture_key, undo_key=self.undo_key,
            monitor=self.selected_monitor(), note_slides=self.note_var.get(),
            skip_duplicates=self.dup_var.get(),
            beep=self.beep_var.get())
        for key, variable in self.format_vars.items():
            self.settings["fmt_" + key] = variable.get()
        save_settings(self.settings)

    # ---- 4d. Choosing what part of the screen to capture --------------------
    def on_monitor_changed(self, _choice):
        self.use_full_screen()          # a region belongs to one monitor, so reset it

    def use_full_screen(self):
        self.region = None
        self.area_label.configure(text="Capture area: full screen")

    def pick_region(self):
        """Show a see-through screen so you can drag a box around the lecture."""
        with MSS() as sct:
            monitor = sct.monitors[self.selected_monitor()]

        self.attributes("-alpha", 0)    # hide our own window for a moment
        self.update()
        time.sleep(0.2)

        # A borderless, see-through window that covers the whole monitor
        cover = tk.Toplevel(self)
        cover.overrideredirect(True)
        cover.geometry(f"{monitor['width']}x{monitor['height']}+{monitor['left']}+{monitor['top']}")
        cover.attributes("-alpha", 0.35, "-topmost", True)
        cover.configure(bg="black")
        canvas = tk.Canvas(cover, bg="black", highlightthickness=0, cursor="cross")
        canvas.pack(fill="both", expand=True)
        cover.update()
        canvas.create_text(cover.winfo_width() // 2, 60, fill="white", font=("Segoe UI", 18),
                           text="Drag a box over the lecture area  -  press Esc to cancel")
        cover.focus_force()

        start = {"x": 0, "y": 0, "box": None}    # where the mouse was pressed

        def close_cover():
            cover.destroy()
            self.attributes("-alpha", 1)         # show our window again

        def mouse_down(event):
            start["x"], start["y"] = event.x, event.y
            start["box"] = canvas.create_rectangle(event.x, event.y, event.x, event.y,
                                                   outline="#00ff88", width=3)

        def mouse_move(event):
            canvas.coords(start["box"], start["x"], start["y"], event.x, event.y)

        def mouse_up(event):
            # Convert the dragged box into real screen pixels
            scale_x = monitor["width"] / max(cover.winfo_width(), 1)
            scale_y = monitor["height"] / max(cover.winfo_height(), 1)
            x1, x2 = sorted((start["x"], event.x))
            y1, y2 = sorted((start["y"], event.y))
            width, height = int((x2 - x1) * scale_x), int((y2 - y1) * scale_y)
            close_cover()
            if width < 40 or height < 40:
                self.set_status("Box was too small - ignored", ORANGE)
                return
            self.region = {"left": monitor["left"] + int(x1 * scale_x),
                           "top": monitor["top"] + int(y1 * scale_y),
                           "width": width, "height": height}
            self.area_label.configure(text=f"Capture area: {width} x {height} box")

        canvas.bind("<ButtonPress-1>", mouse_down)
        canvas.bind("<B1-Motion>", mouse_move)
        canvas.bind("<ButtonRelease-1>", mouse_up)
        cover.bind("<Escape>", lambda event: close_cover())

    # ---- 4e. Choosing the shortcut keys -------------------------------------
    def record_key(self, which):
        """Click a key button, then press the keys you want to use."""
        if self.listening_for_key:
            return
        self.listening_for_key = True
        button = self.capture_key_button if which == "capture" else self.undo_key_button
        button.configure(text="Press keys...")
        self.set_status("Press your new shortcut (Esc to cancel)", CYAN)

        # Waiting for a key press would freeze the window, so we wait inside a
        # "thread" - a helper that runs alongside the window.
        def wait_for_keys():
            try:
                combination = keyboard.read_hotkey(suppress=False)   # e.g. "ctrl+alt+s"
            except Exception:
                combination = None
            self.after(0, lambda: self.finish_record_key(which, combination))

        threading.Thread(target=wait_for_keys, daemon=True).start()

    def finish_record_key(self, which, combination):
        self.listening_for_key = False
        if combination and combination != "esc":
            if which == "capture":
                self.capture_key = combination
            else:
                self.undo_key = combination
            self.set_status("Now click Save keys to use them", CYAN)
        else:
            self.set_status("Cancelled")
        self.capture_key_button.configure(text=self.capture_key)
        self.undo_key_button.configure(text=self.undo_key)

    def apply_keys(self):
        """Check the new shortcuts, switch to them and remember them."""
        if self.capture_key == self.undo_key:
            self.set_status("Both shortcuts are the same - pick different ones", PINK)
            return
        try:
            keyboard.parse_hotkey(self.capture_key)    # raises an error if invalid
            keyboard.parse_hotkey(self.undo_key)
        except Exception as error:
            self.set_status(f"Invalid shortcut: {error}", PINK)
            return
        if self.lecture_running:                       # switch immediately mid-lecture
            self.unregister_hotkeys()
            if not self.register_hotkeys():
                # new keys failed, so go back to the saved ones
                self.capture_key = self.settings["capture_key"]
                self.undo_key = self.settings["undo_key"]
                self.register_hotkeys()
                return
        self.capture_button.configure(text=f"Capture ({self.capture_key})")
        self.undo_button.configure(text=f"Delete ({self.undo_key})")
        self.remember_settings()
        self.set_status("Shortcuts saved", GREEN)

    def register_hotkeys(self):
        """Tell the computer to run our functions when the shortcuts are pressed."""
        try:
            # self.after(0, ...) hands the job over to the main window safely
            self.hotkey_handles = [
                keyboard.add_hotkey(self.capture_key, lambda: self.after(0, self.take_screenshot)),
                keyboard.add_hotkey(self.undo_key, lambda: self.after(0, self.delete_last)),
            ]
            return True
        except Exception as error:
            self.unregister_hotkeys()
            self.set_status(f"Shortcut problem: {error}", PINK)
            return False

    def unregister_hotkeys(self):
        for handle in self.hotkey_handles:
            try:
                keyboard.remove_hotkey(handle)
            except Exception:
                pass
        self.hotkey_handles = []

    # ---- 4f. The lecture itself: start, capture, delete, stop ---------------
    def start_lecture(self):
        if not self.register_hotkeys():
            return
        self.remember_settings()
        # A new temporary folder named after the current time, e.g. 20261011_113000
        self.session_folder = TEMP_FOLDER / datetime.now().strftime("%Y%m%d_%H%M%S")
        self.session_folder.mkdir(parents=True, exist_ok=True)
        self.screenshots = []
        self.last_image = None
        self.lecture_running = True
        self.start_button.configure(state="disabled")
        for button in (self.capture_button, self.undo_button, self.stop_button):
            button.configure(state="normal")
        self.update_counter()
        self.set_status("Recording... press your capture key", GREEN)

    def take_screenshot(self):
        # Ignore the key press if no lecture is running or we are choosing a new key
        if not self.lecture_running or self.listening_for_key:
            return

        self.attributes("-alpha", 0)       # hide our window so it is not in the picture
        self.update()
        time.sleep(0.12)
        try:
            with MSS() as sct:
                area = self.region or sct.monitors[self.selected_monitor()]
                raw = sct.grab(area)
            image = Image.frombytes("RGB", raw.size, raw.bgra, "raw", "BGRX")
        except Exception as error:
            self.set_status(f"Screenshot failed: {error}", PINK)
            return
        finally:
            self.attributes("-alpha", 1)   # always show our window again

        # Optionally skip a screenshot that looks the same as the previous one
        if self.dup_var.get() and self.last_image is not None:
            if images_look_same(image, self.last_image):
                self.set_status("Skipped - looks the same as the last one", ORANGE)
                return

        # Name the file with the next number: 001.png, 002.png, ... (this keeps the order)
        file_path = self.session_folder / f"{len(self.screenshots) + 1:03d}.png"
        image.save(file_path)
        self.screenshots.append(file_path)
        self.last_image = image
        self.update_counter()
        self.set_status(f"Saved page {len(self.screenshots)}", GREEN)
        if self.beep_var.get() and winsound:
            winsound.MessageBeep()

    def delete_last(self):
        if self.lecture_running and self.screenshots and not self.listening_for_key:
            last_file = self.screenshots.pop()      # remove from the list...
            last_file.unlink(missing_ok=True)       # ...and delete the file
            self.last_image = None
            self.update_counter()
            self.set_status("Deleted the last screenshot", ORANGE)

    def stop_and_export(self):
        if not self.screenshots:
            self.finish_lecture()
            self.set_status("No screenshots were taken")
            return

        # Which formats are ticked? We need at least one.
        formats = [key for key, variable in self.format_vars.items() if variable.get()]
        if not formats:
            self.set_status("Tick at least one export format first", PINK)
            return
        if "docx" in formats and Document is None:
            self.set_status("Word needs:  python -m pip install python-docx", PINK)
            return

        # Ask where to save. The lecture name is suggested as the file name,
        # and you can change it in the window.
        suggested = clean_file_name(self.name_box.get()) or "Lecture_" + datetime.now().strftime("%Y-%m-%d")
        first_ending = FORMATS[formats[0]][1]
        chosen = filedialog.asksaveasfilename(
            parent=self, title="Save lecture notes (every ticked format uses this name)",
            initialdir=self.settings["last_folder"],
            initialfile=suggested + first_ending, defaultextension=first_ending,
            filetypes=[(FORMATS[key][0], "*" + FORMATS[key][1]) for key in formats])
        if not chosen:      # you pressed Cancel: keep recording, nothing is lost
            self.set_status("Export cancelled - still recording", ORANGE)
            return

        # Remove the file ending the dialog added; we add the right one per format
        base = Path(chosen).with_suffix("")
        self.settings["last_folder"] = str(base.parent)
        self.set_status("Building your files...", CYAN)
        self.update()

        notes = self.note_var.get()
        saved_files = []
        try:
            for key in formats:
                target = base.with_name(base.name + FORMATS[key][1])
                if key == "pptx":
                    build_pptx(self.screenshots, target, notes)
                elif key == "pdf":
                    build_pdf(self.screenshots, target, notes)
                elif key == "docx":
                    build_docx(self.screenshots, target, notes)
                elif key == "zip":
                    build_zip(self.screenshots, target)
                saved_files.append(target)
        except Exception as error:
            # Something went wrong: keep the screenshots so nothing is lost
            self.set_status("Export failed - screenshots were kept", PINK)
            messagebox.showerror("Export failed",
                                 f"{error}\n\nYour screenshots are still in:\n{self.session_folder}")
            return

        page_count = len(self.screenshots)
        self.finish_lecture(delete_images=True)     # success: remove temporary screenshots
        self.remember_settings()
        self.set_status(f"Saved {page_count} pages as {len(saved_files)} file(s)", GREEN)
        try:
            subprocess.Popen(["explorer", "/select,", str(saved_files[0])])   # show it in Explorer
        except Exception:
            pass

    def finish_lecture(self, delete_images=False):
        """Reset everything back to the 'ready' state."""
        self.lecture_running = False
        self.unregister_hotkeys()
        if delete_images and self.session_folder:
            shutil.rmtree(self.session_folder, ignore_errors=True)
        self.screenshots = []
        self.last_image = None
        self.update_counter()
        self.start_button.configure(state="normal")
        for button in (self.capture_button, self.undo_button, self.stop_button):
            button.configure(state="disabled")

    # ---- 4g. Closing the app -------------------------------------------------
    def on_close(self):
        if self.lecture_running and self.screenshots:
            sure = messagebox.askyesno(
                "Quit?", f"{len(self.screenshots)} screenshots are not exported yet.\n"
                         "Quit and delete them?")
            if not sure:
                return
            shutil.rmtree(self.session_folder, ignore_errors=True)
        self.remember_settings()
        keyboard.unhook_all()
        self.destroy()


# ---------------------------------------------------------------------------
# 5. START THE APP - Python begins here when you run the file
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    App().mainloop()
