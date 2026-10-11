"""
Lecture Capture  -  screenshot a lecture with a hotkey, export an ordered PPTX.

Install:   python -m pip install customtkinter mss keyboard python-pptx pillow
Run:       python lecture_capture.py
Make .exe: python -m pip install pyinstaller
           pyinstaller --onefile --noconsole lecture_capture.py

Screenshots are kept in a temporary folder while you record and are DELETED
after the PPTX is saved successfully.
"""
import json
import os
import shutil
import subprocess
import threading
import time
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox

import customtkinter as ctk
import keyboard
import mss
from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
from pptx.util import Inches

try:
    import winsound
except ImportError:  # not on Windows
    winsound = None

MSS = getattr(mss, "MSS", None) or mss.mss  # new name, falls back to old

APP_DIR = Path.home() / "Documents" / "LectureCaptures"
TEMP_DIR = APP_DIR / "_in_progress"
SETTINGS_FILE = APP_DIR / "settings.json"
DEFAULTS = {
    "capture_key": "f8",
    "undo_key": "f9",
    "monitor": 1,
    "note_slides": False,
    "save_pdf": False,
    "skip_duplicates": False,
    "sound": True,
    "last_dir": str(Path.home() / "Documents"),
}
SLIDE_W, SLIDE_H = 13.333, 7.5  # inches (16:9)


# --------------------------------------------------------------- settings ----
def load_settings():
    s = dict(DEFAULTS)
    try:
        s.update(json.loads(SETTINGS_FILE.read_text(encoding="utf-8")))
    except Exception:
        pass
    return s


def save_settings(s):
    try:
        APP_DIR.mkdir(parents=True, exist_ok=True)
        SETTINGS_FILE.write_text(json.dumps(s, indent=2), encoding="utf-8")
    except Exception:
        pass


# ----------------------------------------------------------------- export ----
def build_pptx(image_paths, out_path, note_slides=False):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    for path in image_paths:
        slide = prs.slides.add_slide(blank)
        with Image.open(path) as im:
            w, h = im.size
        scale = min(SLIDE_W / w, SLIDE_H / h)
        pw, ph = w * scale, h * scale
        slide.shapes.add_picture(
            str(path), Inches((SLIDE_W - pw) / 2), Inches((SLIDE_H - ph) / 2),
            Inches(pw), Inches(ph))
        if note_slides:
            prs.slides.add_slide(blank)
    prs.save(str(out_path))


def build_pdf(image_paths, out_path):
    imgs = [Image.open(p).convert("RGB") for p in image_paths]
    imgs[0].save(str(out_path), save_all=True, append_images=imgs[1:])


def safe_name(text):
    return "".join(c for c in text if c.isalnum() or c in " -_").strip()


def looks_same(a, b, threshold=0.5):
    """True if two PIL images are nearly identical (small grayscale compare)."""
    a = a.convert("L").resize((96, 54))
    b = b.convert("L").resize((96, 54))
    return ImageStat.Stat(ImageChops.difference(a, b)).mean[0] < threshold


# -------------------------------------------------------------------- app ----
class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Lecture Capture")
        self.geometry("360x560")
        self.attributes("-topmost", True)
        ctk.set_appearance_mode("dark")

        self.cfg = load_settings()
        self.cap_key = self.cfg['capture_key']
        self.undo_key = self.cfg['undo_key']
        self.rec_active = False
        self.active = False
        self.shots = []
        self.folder = None
        self.region = None          # None = whole monitor
        self.last_img = None
        self.handles = []

        pad = dict(fill="x", padx=16)
        ctk.CTkLabel(self, text="Lecture Capture",
                     font=ctk.CTkFont(size=20, weight="bold")).pack(pady=(12, 4))

        self.name_entry = ctk.CTkEntry(self, placeholder_text="Lecture name (used as file name)")
        self.name_entry.pack(pady=4, **pad)

        with MSS() as sct:
            n = len(sct.monitors) - 1
        self.monitor_menu = ctk.CTkOptionMenu(
            self, values=[f"Monitor {i}" for i in range(1, n + 1)],
            command=lambda _: self.reset_region())
        self.monitor_menu.set(f"Monitor {min(self.cfg['monitor'], n)}")
        self.monitor_menu.pack(pady=4, **pad)

        row = ctk.CTkFrame(self, fg_color="transparent")
        row.pack(pady=2, **pad)
        ctk.CTkButton(row, text="Select region", width=150,
                      command=self.pick_region).pack(side="left", padx=(0, 6))
        ctk.CTkButton(row, text="Full screen", width=150, fg_color="#555",
                      command=self.reset_region).pack(side="left")
        self.region_label = ctk.CTkLabel(self, text="Capture area: full screen",
                                         text_color="gray")
        self.region_label.pack()

        keys = ctk.CTkFrame(self, fg_color="transparent")
        keys.pack(pady=4, **pad)
        ctk.CTkLabel(keys, text="Capture key").grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(keys, text="Delete-last key").grid(row=1, column=0, sticky="w")
        self.cap_key_btn = ctk.CTkButton(
            keys, text=self.cap_key, width=130, fg_color="#333",
            command=lambda: self.record_key("cap"))
        self.cap_key_btn.grid(row=0, column=1, padx=8, pady=2)
        self.undo_key_btn = ctk.CTkButton(
            keys, text=self.undo_key, width=130, fg_color="#333",
            command=lambda: self.record_key("undo"))
        self.undo_key_btn.grid(row=1, column=1, padx=8, pady=2)
        ctk.CTkButton(keys, text="Save keys", width=64,
                      command=self.apply_keys).grid(row=0, column=2, rowspan=2)

        self.note_var = ctk.BooleanVar(value=self.cfg["note_slides"])
        self.pdf_var = ctk.BooleanVar(value=self.cfg["save_pdf"])
        self.dup_var = ctk.BooleanVar(value=self.cfg["skip_duplicates"])
        self.snd_var = ctk.BooleanVar(value=self.cfg["sound"])
        for text, var in (("Blank note slide after each screenshot", self.note_var),
                          ("Also save a PDF", self.pdf_var),
                          ("Skip near-identical captures", self.dup_var),
                          ("Beep on capture", self.snd_var)):
            ctk.CTkCheckBox(self, text=text, variable=var).pack(anchor="w", padx=16, pady=2)

        self.counter = ctk.CTkLabel(self, text="0 pages",
                                    font=ctk.CTkFont(size=26, weight="bold"))
        self.counter.pack(pady=6)

        self.start_btn = ctk.CTkButton(self, text="Start Lecture", command=self.start)
        self.start_btn.pack(pady=3, **pad)
        self.shot_btn = ctk.CTkButton(self, text=f"Capture ({self.cap_key})", command=self.capture,
                                      state="disabled")
        self.shot_btn.pack(pady=3, **pad)
        self.undo_btn = ctk.CTkButton(self, text=f"Delete last ({self.undo_key})", command=self.undo,
                                      state="disabled", fg_color="#555")
        self.undo_btn.pack(pady=3, **pad)
        self.stop_btn = ctk.CTkButton(self, text="Stop & Export", command=self.stop,
                                      state="disabled", fg_color="#b33")
        self.stop_btn.pack(pady=3, **pad)

        self.status = ctk.CTkLabel(self, text="Ready", text_color="gray")
        self.status.pack(pady=4)
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    # ------------------------------------------------------------ helpers ---
    def set_status(self, text):
        self.status.configure(text=text)

    def update_counter(self):
        self.counter.configure(text=f"{len(self.shots)} pages")

    def monitor_index(self):
        return int(self.monitor_menu.get().split()[-1])

    def remember_settings(self):
        self.cfg.update(
            capture_key=self.cap_key, undo_key=self.undo_key,
            monitor=self.monitor_index(), note_slides=self.note_var.get(),
            save_pdf=self.pdf_var.get(), skip_duplicates=self.dup_var.get(),
            sound=self.snd_var.get())
        save_settings(self.cfg)

    # ------------------------------------------------------------- region ---
    def reset_region(self):
        self.region = None
        self.region_label.configure(text="Capture area: full screen")

    def pick_region(self):
        with MSS() as sct:
            mon = sct.monitors[self.monitor_index()]
        self.attributes("-alpha", 0)
        self.update()
        time.sleep(0.2)

        top = tk.Toplevel(self)
        top.overrideredirect(True)
        top.geometry(f"{mon['width']}x{mon['height']}+{mon['left']}+{mon['top']}")
        top.attributes("-alpha", 0.35, "-topmost", True)
        top.configure(bg="black")
        canvas = tk.Canvas(top, bg="black", highlightthickness=0, cursor="cross")
        canvas.pack(fill="both", expand=True)
        top.update()
        canvas.create_text(top.winfo_width() // 2, 60, fill="white",
                           font=("Segoe UI", 18),
                           text="Drag over the lecture area  -  Esc to cancel")
        top.focus_force()
        st = {"x": 0, "y": 0, "rect": None}

        def close():
            top.destroy()
            self.attributes("-alpha", 1)

        def press(e):
            st["x"], st["y"] = e.x, e.y
            st["rect"] = canvas.create_rectangle(e.x, e.y, e.x, e.y,
                                                 outline="#00ff88", width=3)

        def drag(e):
            canvas.coords(st["rect"], st["x"], st["y"], e.x, e.y)

        def release(e):
            sx = mon["width"] / max(top.winfo_width(), 1)
            sy = mon["height"] / max(top.winfo_height(), 1)
            x0, x1 = sorted((st["x"], e.x))
            y0, y1 = sorted((st["y"], e.y))
            w, h = int((x1 - x0) * sx), int((y1 - y0) * sy)
            close()
            if w < 40 or h < 40:
                self.set_status("Region too small - ignored")
                return
            self.region = {"left": mon["left"] + int(x0 * sx),
                           "top": mon["top"] + int(y0 * sy),
                           "width": w, "height": h}
            self.region_label.configure(text=f"Capture area: {w} x {h} region")

        canvas.bind("<ButtonPress-1>", press)
        canvas.bind("<B1-Motion>", drag)
        canvas.bind("<ButtonRelease-1>", release)
        top.bind("<Escape>", lambda e: close())

    # ------------------------------------------------------ session control --
    def register_hotkeys(self):
        cap, und = self.cap_key.strip().lower(), self.undo_key.strip().lower()
        try:
            self.handles = [
                keyboard.add_hotkey(cap, lambda: self.after(0, self.capture)),
                keyboard.add_hotkey(und, lambda: self.after(0, self.undo)),
            ]
        except Exception as e:
            self.unregister_hotkeys()
            self.set_status(f"Bad hotkey: {e}")
            return False
        self.shot_btn.configure(text=f"Capture ({cap})")
        self.undo_btn.configure(text=f"Delete last ({und})")
        return True

    def unregister_hotkeys(self):
        for h in self.handles:
            try:
                keyboard.remove_hotkey(h)
            except Exception:
                pass
        self.handles = []

    # ------------------------------------------------------------- keys ----
    def refresh_key_buttons(self):
        self.cap_key_btn.configure(text=self.cap_key)
        self.undo_key_btn.configure(text=self.undo_key)

    def record_key(self, which):
        """Click the button, then press the key combination you want."""
        if self.rec_active:
            return
        self.rec_active = True
        (self.cap_key_btn if which == "cap" else self.undo_key_btn).configure(
            text="Press keys...")
        self.set_status("Press the key combination (Esc = cancel)")

        def worker():
            try:
                combo = keyboard.read_hotkey(suppress=False)
            except Exception:
                combo = None
            self.after(0, lambda: self.finish_record(which, combo))

        threading.Thread(target=worker, daemon=True).start()

    def finish_record(self, which, combo):
        self.rec_active = False
        if combo and combo != "esc":
            if which == "cap":
                self.cap_key = combo
            else:
                self.undo_key = combo
            self.set_status("Keys changed - click Save keys")
        else:
            self.set_status("Cancelled")
        self.refresh_key_buttons()

    def apply_keys(self):
        cap, und = self.cap_key.strip().lower(), self.undo_key.strip().lower()
        if cap == und:
            self.set_status("Both keys are the same - pick different keys")
            return
        try:
            keyboard.parse_hotkey(cap)
            keyboard.parse_hotkey(und)
        except Exception as e:
            self.set_status(f"Invalid key: {e}")
            return
        if self.active:                       # apply immediately mid-lecture
            self.unregister_hotkeys()
            if not self.register_hotkeys():   # failed -> go back to saved keys
                self.cap_key, self.undo_key = self.cfg["capture_key"], self.cfg["undo_key"]
                self.register_hotkeys()
                self.refresh_key_buttons()
                return
        self.shot_btn.configure(text=f"Capture ({cap})")
        self.undo_btn.configure(text=f"Delete last ({und})")
        self.remember_settings()
        self.set_status("Keys saved")

    def start(self):
        if not self.register_hotkeys():
            return
        self.remember_settings()
        self.folder = TEMP_DIR / datetime.now().strftime("%Y%m%d_%H%M%S")
        self.folder.mkdir(parents=True, exist_ok=True)
        self.shots, self.last_img, self.active = [], None, True
        self.start_btn.configure(state="disabled")
        for b in (self.shot_btn, self.undo_btn, self.stop_btn):
            b.configure(state="normal")
        self.update_counter()
        self.set_status("Recording...")

    def capture(self):
        if not self.active or self.rec_active:
            return
        self.attributes("-alpha", 0)       # keep this window out of the shot
        self.update()
        time.sleep(0.12)
        try:
            with MSS() as sct:
                box = self.region or sct.monitors[self.monitor_index()]
                raw = sct.grab(box)
            img = Image.frombytes("RGB", raw.size, raw.bgra, "raw", "BGRX")
        except Exception as e:
            self.set_status(f"Capture failed: {e}")
            return
        finally:
            self.attributes("-alpha", 1)

        if self.dup_var.get() and self.last_img is not None and looks_same(img, self.last_img):
            self.set_status("Skipped - same as last")
            return
        path = self.folder / f"{len(self.shots) + 1:03d}.png"
        img.save(path)
        self.shots.append(path)
        self.last_img = img
        self.update_counter()
        self.set_status(f"Saved page {len(self.shots)}")
        if self.snd_var.get() and winsound:
            winsound.MessageBeep()

    def undo(self):
        if self.active and self.shots and not self.rec_active:
            self.shots.pop().unlink(missing_ok=True)
            self.last_img = None
            self.update_counter()
            self.set_status("Deleted last capture")

    def stop(self):
        if not self.shots:
            self.end_session()
            self.set_status("No screenshots taken")
            return

        name = safe_name(self.name_entry.get()) or "Lecture_" + datetime.now().strftime("%Y-%m-%d")
        chosen = filedialog.asksaveasfilename(
            parent=self, title="Save lecture notes", initialdir=self.cfg["last_dir"],
            initialfile=name + ".pptx", defaultextension=".pptx",
            filetypes=[("PowerPoint", "*.pptx")])
        if not chosen:                      # cancelled -> keep recording
            self.set_status("Export cancelled - still recording")
            return

        out = Path(chosen)
        self.cfg["last_dir"] = str(out.parent)
        self.set_status("Building PPTX...")
        self.update()
        try:
            build_pptx(self.shots, out, self.note_var.get())
            if self.pdf_var.get():
                build_pdf(self.shots, out.with_suffix(".pdf"))
        except Exception as e:              # keep screenshots if export fails
            self.set_status("Export failed - screenshots kept")
            messagebox.showerror("Export failed",
                                 f"{e}\n\nScreenshots are kept in:\n{self.folder}")
            return

        count = len(self.shots)
        self.end_session(delete_images=True)
        self.set_status(f"Saved {count} slides")
        self.remember_settings()
        try:
            subprocess.Popen(["explorer", "/select,", str(out)])
        except Exception:
            pass

    def end_session(self, delete_images=False):
        self.active = False
        self.unregister_hotkeys()
        if delete_images and self.folder:
            shutil.rmtree(self.folder, ignore_errors=True)
        self.shots, self.last_img = [], None
        self.update_counter()
        self.start_btn.configure(state="normal")
        for b in (self.shot_btn, self.undo_btn, self.stop_btn):
            b.configure(state="disabled")

    def on_close(self):
        if self.active and self.shots:
            if not messagebox.askyesno(
                    "Quit?", f"{len(self.shots)} screenshots are not exported yet.\n"
                             "Quit and delete them?"):
                return
            shutil.rmtree(self.folder, ignore_errors=True)
        self.remember_settings()
        keyboard.unhook_all()
        self.destroy()


if __name__ == "__main__":
    App().mainloop()
